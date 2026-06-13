#!/usr/bin/env python3
"""Structural QA gate for a finished Keynote Deck Builder deck (Phase 3).

Checks (HARD -- non-zero exit, so this can gate a build):
  * slide count matches --expect (when given),
  * every slide carries speaker notes (the spoken beat lives in the notes, not on the
    slide), and
  * NO leftover icon-font glyphs -- the exact failure mode the visual rebuild exists to
    fix (keynote decks use vector motifs, never icon fonts).

Checks (SOFT -- reported, but only fail the run under --strict):
  * notes word count outside the band (default 6-80): the floor flags a slide whose notes
    look too thin to be a real spoken beat; the ceiling flags a beat heavy enough it should
    probably be split, and
  * suspicious snake_case tokens that look like dropped icon-font ligatures.

These decks are spoken talks: short on-slide text, no inline citations, and often no numbers
by the source's discipline. So there is deliberately NO citation/number/fidelity check here
-- that belongs to research-deck-builder, not this skill.

The title and closing slides are exempt from the word band by default (house style keeps
them short); override with --band-exempt.

Usage:
  python3 verify_keynote.py --deck Talk_KEYNOTE.pptx --expect 32
  python3 verify_keynote.py --deck Talk_KEYNOTE.pptx --expect 32 --strict
  python3 verify_keynote.py --deck Talk_KEYNOTE.pptx --no-require-notes
"""
import argparse
import re
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Icon fonts that render as broken text where they are not installed. Any run using one
# of these is a hard failure: rebuild that slide with vector shapes instead (see SKILL.md).
ICON_FONTS = re.compile(r"material\s*(symbols|icons)|font\s*awesome|glyphicons", re.I)

# Material-Symbols ligatures look like normal snake_case words once the font drops out
# (e.g. 'check_circle', 'arrow_forward'). Flag short all-lowercase underscore tokens that
# sit alone in a run -- a soft signal, since legitimate code samples also use snake_case.
LIGATURE = re.compile(r"^[a-z]{2,}(?:_[a-z]{2,}){1,3}$")


def walk(shapes):
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(s.shapes)
        else:
            yield s


def parse_exempt(spec, n_slides):
    out = set()
    for tok in (spec or "").split(","):
        tok = tok.strip().lower()
        if not tok:
            continue
        if tok == "first":
            out.add(1)
        elif tok == "last":
            out.add(n_slides)
        else:
            out.add(int(tok))
    return out


def verify(args):
    p = Presentation(args.deck)
    slides = list(p.slides)
    n = len(slides)
    band_exempt = parse_exempt(args.band_exempt, n)

    hard, soft = [], []

    # --- slide count ---
    if args.expect is not None and n != args.expect:
        hard.append(f"slide count is {n}, expected {args.expect}")

    for i, s in enumerate(slides, 1):
        # collect text + fonts on this slide
        texts, fonts = [], set()
        for sh in walk(s.shapes):
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text:
                        texts.append(r.text)
                    if r.font.name:
                        fonts.add(r.font.name)

        # --- icon-font (hard) + ligature artifacts (soft) ---
        bad_fonts = [f for f in fonts if ICON_FONTS.search(f or "")]
        if bad_fonts:
            hard.append(f"S{i}: icon font still in use -> {', '.join(sorted(bad_fonts))}")
        for tok in texts:
            t = tok.strip()
            if LIGATURE.match(t) and "Consolas" not in fonts:
                soft.append(f"S{i}: possible icon-ligature artifact text '{t}'")

        # --- notes (hard if required; soft word band) ---
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        wc = len(notes.split())
        if args.require_notes and not notes:
            hard.append(f"S{i}: no speaker notes")
        elif notes and i not in band_exempt and not (args.min_words <= wc <= args.max_words):
            soft.append(f"S{i}: notes {wc} words (band {args.min_words}-{args.max_words})")

    # --- report ---
    print(f"Deck: {args.deck}")
    print(f"Slides: {n}" + (f" (expected {args.expect})" if args.expect is not None else ""))
    print(f"HARD failures: {len(hard)} | SOFT warnings: {len(soft)}\n")
    for h in hard:
        print(f"  [FAIL] {h}")
    for w in soft:
        print(f"  [warn] {w}")
    if not hard and not soft:
        print("  All checks passed.")

    failed = bool(hard) or (args.strict and bool(soft))
    print("\nRESULT:", "FAIL" if failed else "PASS")
    return 1 if failed else 0


def main():
    ap = argparse.ArgumentParser(description="Structural QA gate for a keynote deck (.pptx).")
    ap.add_argument("--deck", required=True, help="deck to check")
    ap.add_argument("--expect", type=int, help="expected slide count")
    ap.add_argument("--min-words", type=int, default=6, help="notes word floor (default 6)")
    ap.add_argument("--max-words", type=int, default=80, help="notes word ceiling (default 80)")
    ap.add_argument("--band-exempt", default="first,last",
                    help="slides exempt from the word band (default: first,last)")
    ap.add_argument("--require-notes", dest="require_notes", action="store_true", default=True,
                    help="fail if any slide lacks notes (default)")
    ap.add_argument("--no-require-notes", dest="require_notes", action="store_false",
                    help="don't require notes")
    ap.add_argument("--strict", action="store_true", help="treat SOFT warnings as failures")
    sys.exit(verify(ap.parse_args()))


if __name__ == "__main__":
    main()
