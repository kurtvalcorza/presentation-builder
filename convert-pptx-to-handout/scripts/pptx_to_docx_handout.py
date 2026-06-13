#!/usr/bin/env python3
"""
Convert PowerPoint to DOCX Handouts or Markdown Transcripts
Extracts slides and speaker notes into a formatted Word document or Markdown file
"""

import sys
import os
import re

try:
    from pptx import Presentation
except ModuleNotFoundError:
    print("Error: python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)

def clean_text(text):
    """Remove XML-incompatible characters from text"""
    if not text:
        return ""
    # Remove control characters except newline, carriage return, and tab
    text = ''.join(char for char in text if ord(char) >= 32 or char in '\n\r\t')
    # Remove any remaining problematic characters
    text = re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F-\x9F]', '', text)
    return text.strip()

def extract_to_docx(prs, docx_path):
    """Save presentation slides and notes to a DOCX handout"""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    title = doc.add_heading('Presentation Handout', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for idx, slide in enumerate(prs.slides, 1):
        # Slide number heading
        doc.add_heading(f'Slide {idx}', level=1)
        
        # Slide title
        if slide.shapes.title:
            title_text = clean_text(slide.shapes.title.text)
            if title_text:
                p = doc.add_paragraph()
                run = p.add_run(f"Title: {title_text}")
                run.bold = True
                run.font.size = Pt(12)
        
        # Slide content
        text_content = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                if shape != slide.shapes.title:
                    cleaned = clean_text(shape.text)
                    if cleaned:
                        text_content.append(cleaned)
        
        if text_content:
            doc.add_heading('Slide Content:', level=2)
            for text in text_content:
                doc.add_paragraph(text, style='List Bullet')
        
        # Speaker notes
        if slide.has_notes_slide and slide.notes_slide:
            notes_text = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes_text:
                doc.add_heading('Speaker Notes:', level=2)
                doc.add_paragraph(notes_text)
                
        # Page break
        if idx < len(prs.slides):
            doc.add_page_break()
            
    doc.save(docx_path)

def extract_to_md(prs, md_path, pptx_filename):
    """Save presentation slides and notes directly to a Markdown note"""
    md_content = []
    presentation_title = os.path.splitext(pptx_filename)[0]
    
    # Metadata Frontmatter
    md_content.append("---")
    md_content.append(f'title: "{presentation_title}"')
    md_content.append("type: presentation-transcript")
    md_content.append("status: processed")
    md_content.append("tags: [presentation, transcript, reference]")
    md_content.append("links:")
    md_content.append(f'  - "[[{pptx_filename}]]"')
    md_content.append("---")
    md_content.append("")
    md_content.append(f"# {presentation_title}")
    md_content.append(f"*Extracted from presentation source: `{pptx_filename}`*\n")
    md_content.append("---")
    md_content.append("")
    
    for idx, slide in enumerate(prs.slides, 1):
        md_content.append(f"## Slide {idx}")
        
        # Slide Title
        if slide.shapes.title:
            slide_title = clean_text(slide.shapes.title.text)
            if slide_title:
                md_content.append(f"### {slide_title}\n")
        
        # Slide Content
        slide_texts = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if hasattr(shape, "text") and shape.text.strip():
                cleaned = clean_text(shape.text)
                if cleaned:
                    paragraphs = [p.strip() for p in cleaned.split('\n') if p.strip()]
                    for p in paragraphs:
                        slide_texts.append(f"- {p}")
                        
        if slide_texts:
            md_content.append("**Slide Content:**")
            md_content.extend(slide_texts)
            md_content.append("")
        else:
            md_content.append("*[No text content on slide]*\n")
            
        # Speaker Notes
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide:
            notes_text = clean_text(slide.notes_slide.notes_text_frame.text)
            
        if notes_text:
            md_content.append("**Speaker Notes:**")
            notes_lines = notes_text.split('\n')
            for line in notes_lines:
                if line.strip():
                    md_content.append(f"> {line.strip()}")
            md_content.append("")
        else:
            md_content.append("**Speaker Notes:**\n*[No speaker notes]*\n")
            
        md_content.append("---\n")
        
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md_content))

def main():
    args = [a for a in sys.argv[1:] if a != "--force"]
    force = "--force" in sys.argv
    if len(args) < 1:
        print("Usage: python pptx_to_docx_handout.py <input.pptx> [output_file] [--force]")
        sys.exit(1)
        
    pptx_path = args[0]
    if not os.path.exists(pptx_path):
        print(f"Error: File not found: {pptx_path}")
        sys.exit(1)
    if not pptx_path.lower().endswith(".pptx"):
        print(f"Error: '{pptx_path}' is not a .pptx file (convert .ppt to .pptx first, or check the path).")
        sys.exit(1)
        
    output_path = args[1] if len(args) > 1 else None
    
    if output_path is None:
        # Default to docx in same directory
        output_path = os.path.splitext(pptx_path)[0] + "_handout.docx"

    # Overwrite guard: don't silently clobber an existing handout
    if os.path.exists(output_path) and not force:
        print(f"Error: {output_path} already exists. Re-run with --force to overwrite.")
        sys.exit(1)

    # Ensure the output directory exists
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.isdir(out_dir):
        try:
            os.makedirs(out_dir, exist_ok=True)
        except OSError as e:
            print(f"Error: cannot create output directory '{out_dir}': {e}")
            sys.exit(1)
        
    ext = os.path.splitext(output_path)[1].lower()
    
    print(f"Reading: {pptx_path}")
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"Error: could not open '{pptx_path}' as a PPTX (corrupt, password-protected, or not a valid .pptx): {e}")
        sys.exit(1)
    if len(prs.slides) == 0:
        print("Warning: the presentation has no slides; output will be near-empty.")
    
    if ext == '.md':
        print("Extracting directly to Markdown...")
        extract_to_md(prs, output_path, os.path.basename(pptx_path))
    else:
        # Ensure we have python-docx installed for DOCX mode
        try:
            import docx
        except ImportError:
            print("Error: python-docx is required for DOCX handouts. Install with: pip install python-docx")
            sys.exit(1)
            
        print("Extracting to DOCX handout...")
        extract_to_docx(prs, output_path)
        
    print(f"\nSuccess! Saved to: {output_path}")
    print(f"Total slides processed: {len(prs.slides)}")

if __name__ == "__main__":
    main()
